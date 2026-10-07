from datetime import timedelta
from io import BytesIO
import csv
from textwrap import wrap

from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from tickets.access import visible_tickets_for
from tickets.models import TicketPriority, WorkflowStatus
from core.models import Division
from .forms import ReportDateFilterForm


@login_required
def dashboard(request):
    tickets = visible_tickets_for(request.user)
    full_dashboard = request.user.is_super_admin or request.user.is_director or request.user.is_senior_it
    dashboard_tickets = tickets if full_dashboard else tickets.filter(
        Q(requester=request.user) | Q(current_assignee=request.user)
    ).distinct()
    now = timezone.now()
    open_tickets = dashboard_tickets.exclude(status__kind__in=["RESOLVED", "CLOSED"])
    metrics = {
        "total": dashboard_tickets.count(),
        "open": open_tickets.count(),
        "in_progress": dashboard_tickets.filter(status__code="in-progress").count(),
        "verification": dashboard_tickets.filter(status__code__in=["ready-for-it-verification", "ready-for-verification"]).count(),
        "overdue": open_tickets.filter(deadline__lt=now).count(),
        "completed": dashboard_tickets.filter(status__kind__in=["RESOLVED", "CLOSED"]).count(),
    }
    division_counts = list(
        tickets.values("requesting_division__name", "requesting_division__code")
        .annotate(total=Count("id"), open=Count("id", filter=~Q(status__kind__in=["RESOLVED", "CLOSED"])))
        .order_by("requesting_division__code")
    )
    max_division = max([row["total"] for row in division_counts], default=1)
    for row in division_counts:
        row["percent"] = round(row["total"] / max_division * 100)
    priority_counts = list(tickets.values("priority__name", "priority__color").annotate(total=Count("id")).order_by("-priority__rank"))
    context = {
        "metrics": metrics,
        "full_dashboard": full_dashboard,
        "division_counts": division_counts,
        "priority_counts": priority_counts,
        "recent_tickets": dashboard_tickets[:7],
        "due_soon": open_tickets.filter(deadline__gte=now, deadline__lte=now + timedelta(days=7)).order_by("deadline")[:5],
        "my_tasks": open_tickets.filter(current_assignee=request.user).order_by("deadline")[:5],
        "my_requests": dashboard_tickets.filter(requester=request.user).order_by("-created_at")[:5],
    }
    return render(request, "dashboard/dashboard.html", context)


def _report_queryset(request):
    tickets = visible_tickets_for(request.user)
    division = request.GET.get("division", "")
    status = request.GET.get("status", "")
    priority = request.GET.get("priority", "")
    dates = ReportDateFilterForm(request.GET)
    if not dates.is_valid():
        raise ValidationError("Please enter valid From and To dates, with From on or before To.")
    start = dates.cleaned_data["start"]
    end = dates.cleaned_data["end"]
    if division:
        tickets = tickets.filter(requesting_division__code=division)
    if status:
        tickets = tickets.filter(status__code=status)
    if priority:
        tickets = tickets.filter(priority__code=priority)
    if start:
        tickets = tickets.filter(created_at__date__gte=start)
    if end:
        tickets = tickets.filter(created_at__date__lte=end)
    return tickets


@login_required
def reports(request):
    filter_error = ""
    try:
        tickets = _report_queryset(request)
    except ValidationError as error:
        filter_error = error.messages[0]
        tickets = visible_tickets_for(request.user).none()
    now = timezone.now()
    resolved = tickets.filter(status__kind__in=["RESOLVED", "CLOSED"])
    durations = [
        (completed_at - created_at).total_seconds() / 3600
        for created_at, completed_at in resolved.values_list("created_at", "completed_at") if completed_at
    ]
    context = {
        "filter_error": filter_error,
        "metrics": {
            "total": tickets.count(), "open": tickets.exclude(status__kind__in=["RESOLVED", "CLOSED"]).count(),
            "resolved": resolved.count(), "overdue": tickets.filter(deadline__lt=now).exclude(status__kind__in=["RESOLVED", "CLOSED"]).count(),
            "reopened": tickets.filter(status__code="reopened").count(),
            "avg_hours": round(sum(durations) / len(durations), 1) if durations else 0,
        },
        "status_data": list(tickets.values("status__name", "status__color", "status__kind").annotate(total=Count("id")).order_by("status__sort_order")),
        "officers": list(
            tickets.exclude(current_assignee=None)
            .values("current_assignee_id", "current_assignee__first_name", "current_assignee__last_name", "current_assignee__username")
            .annotate(total=Count("id"), open=Count("id", filter=~Q(status__kind__in=["RESOLVED", "CLOSED"])))
            .order_by("-open", "-total", "current_assignee__username")
        ),
        "report_tickets": tickets.select_related("identified_by", "reviewed_by").order_by("-created_at", "-pk"),
        "division_options": Division.objects.filter(is_active=True), "statuses": WorkflowStatus.objects.filter(is_active=True),
        "priorities": TicketPriority.objects.filter(is_active=True), "filters": request.GET,
    }
    return render(request, "dashboard/reports.html", context, status=400 if filter_error else 200)


def _report_rows(tickets, detailed=False):
    if detailed:
        yield ["Ticket ID", "Title", "Division", "Status", "Sent by", "Assigned to", "Created", "Updated", "Deadline", "Completed"]
    else:
        yield ["Ticket ID", "Title", "Division", "Product", "Category", "Priority", "Status", "Requester", "Assignee", "Created", "Deadline", "Completed"]
    for ticket in tickets.select_related("requesting_division", "product", "category", "priority", "status", "requester", "current_assignee", "identified_by", "reviewed_by"):
        row = [
            ticket.ticket_number, ticket.title, ticket.requesting_division.name, ticket.product.name,
            ticket.category.name, ticket.priority.name, ticket.status.name, ticket.requester.full_name,
            ticket.current_assignee.full_name if ticket.current_assignee else "",
            timezone.localtime(ticket.created_at).strftime("%Y-%m-%d %H:%M"),
            timezone.localtime(ticket.deadline).strftime("%Y-%m-%d %H:%M") if ticket.deadline else "",
            timezone.localtime(ticket.completed_at).strftime("%Y-%m-%d %H:%M") if ticket.completed_at else "",
        ]
        if detailed:
            row = row[:3] + [row[6], row[7], row[8] or "Unassigned",
                row[9], timezone.localtime(ticket.updated_at).strftime("%Y-%m-%d %H:%M"),
                row[10] or "No deadline", row[11] or "Pending"]
        yield row


@login_required
def report_export(request, format):
    try:
        tickets = _report_queryset(request).order_by("-created_at")
    except ValidationError as error:
        return HttpResponse(error.messages[0], status=400, content_type="text/plain")
    rows = list(_report_rows(tickets, detailed=format == "xlsx"))
    filename = f"frc-ticket-report-{timezone.localdate().isoformat()}"
    if format == "csv":
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        writer = csv.writer(response)
        writer.writerows(rows)
        return response
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.page import PageMargins
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Ticket Report"
        for row in rows:
            sheet.append(row)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.data_type = "s"
        widths = [11, 18, 6, 12, 12, 12, 10, 10, 10, 10]
        edge = Side(style="thin", color="D9E1E8")
        border = Border(left=edge, right=edge, top=edge, bottom=edge)
        for column, width in zip(sheet.columns, widths):
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
                cell.font = Font(name="Calibri", size=10)
                if cell.row == 1:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="173D5F")
            # Excel does not reliably auto-fit wrapped rows when printing.
            lines = max(
                sum(max(1, len(wrap(part, width=max(1, int(width) - 2))))
                    for part in str(cell.value or "").split("\n"))
                for cell, width in zip(row, widths)
            )
            sheet.row_dimensions[row[0].row].height = min(409, max(32 if row[0].row == 1 else 60, lines * 15 + 12))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.2, footer=0.2)
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:1"
        sheet.print_area = sheet.dimensions
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return response
    if format == "pdf":
        from .report_pdf import build_ticket_pdf
        response = HttpResponse(build_ticket_pdf(tickets), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
        return response
    return HttpResponse("Unsupported export format", status=400)
